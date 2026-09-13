"""
Excel 导入导出共享逻辑
供 Flask 导出接口与 gzb_output.py 等脚本共用。
"""

import io
import os
import re

import pandas as pd
from openpyxl import load_workbook

from db import query_all, get_connection, get_pinyin_sx

# 导出 Excel 的列名（与模板保持一致）
EXPORT_COLUMNS = ['姓名', '身份证号', '银行账号', '银行地址', '电话']


def read_names(path: str, limit: int = None) -> list:
    """
    读取 Excel 第一列的姓名列表。

    参数:
        path:  Excel 文件路径
        limit: 最多读取的行数，None 表示全部读取
    返回:
        去空白后的姓名列表（保留原顺序，不去重）
    异常:
        FileNotFoundError: 文件不存在时抛出
    """
    if not os.path.exists(path):
        raise FileNotFoundError(f'{path} 文件不存在')

    wb = load_workbook(path, read_only=True)
    try:
        ws = wb.active
        names = []
        for row in ws.iter_rows(max_row=limit, values_only=True):
            if row and row[0] is not None and str(row[0]).strip():
                names.append(str(row[0]).strip())
        # 首行是表头（"姓名"）时自动跳过，避免被当成姓名查询
        if names and names[0] == '姓名':
            names = names[1:]
        return names
    finally:
        wb.close()


def build_employee_dataframe(names: list) -> pd.DataFrame:
    """
    按姓名列表从数据库匹配员工信息，生成导出用 DataFrame。
    一次性取出全表建立索引，避免逐姓名查询（N+1 问题）；
    只匹配在职记录（deleted = 0），数据库中不存在或已删除的姓名导出为空行。
    """
    rows = query_all(
        "SELECT real_name, id_number, bank_account, bank_address, phone "
        "FROM employees WHERE deleted = 0"
    )
    employees_by_name = {row['real_name']: row for row in rows}

    records = []
    for name in names:
        emp = employees_by_name.get(name)
        records.append({
            '姓名': name,
            '身份证号': (emp['id_number'] or '') if emp else '',
            '银行账号': (emp['bank_account'] or '') if emp else '',
            '银行地址': (emp['bank_address'] or '') if emp else '',
            '电话': (emp['phone'] or '') if emp else '',
        })
    return pd.DataFrame(records, columns=EXPORT_COLUMNS)


def build_export_buffer(names: list) -> tuple:
    """
    按姓名列表生成导出 Excel，写入内存缓冲区（不落盘，供 HTTP 直接返回）。

    返回: (BytesIO 缓冲区, 总条数, 数据库匹配到的条数)
    """
    df = build_employee_dataframe(names)
    matched = int((df['身份证号'] != '').sum()) if len(df) else 0
    buf = io.BytesIO()
    df.to_excel(buf, index=False, engine='openpyxl')
    buf.seek(0)
    return buf, len(df), matched


def export_employees(names: list, output_path: str) -> tuple:
    """
    按姓名列表导出员工信息到 Excel 文件（供命令行脚本使用）。

    返回: (总条数, 数据库匹配到的条数)
    """
    buf, total, matched = build_export_buffer(names)
    with open(output_path, 'wb') as f:
        f.write(buf.getvalue())
    return total, matched


# ---------------------------------------------------------------------------
# 导入（Web 上传接口与 gzb_update.py 共用）
# ---------------------------------------------------------------------------

# 导入 Excel 需要的数据库字段
IMPORT_COLUMNS = ['real_name', 'id_number', 'bank_account', 'bank_address', 'phone']

# 表头别名映射：自动识别中英文表头（键为表头文字，值为标准字段名）
HEADER_ALIASES = {
    'real_name': 'real_name', '姓名': 'real_name', 'name': 'real_name',
    'id_number': 'id_number', '身份证号': 'id_number', '身份证': 'id_number',
    'bank_account': 'bank_account', '银行账号': 'bank_account', '银行卡号': 'bank_account',
    'bank_address': 'bank_address', '银行地址': 'bank_address', '开户行': 'bank_address',
    'phone': 'phone', '电话': 'phone', '手机号': 'phone', '手机号码': 'phone',
}

# 导入行状态对应的中文标签（前端展示用）
STATUS_LABELS = {'new': '新增', 'update': '更新', 'restore': '恢复', 'error': '错误'}


def _clean_cell(value) -> str:
    """
    清洗单元格文本：转字符串并去空白；
    数值单元格被 pandas 读出的小数尾巴（如 '13800000000.0'）还原为整数字符串。
    """
    text = str(value).strip()
    if re.fullmatch(r'\d+\.0', text):
        text = text[:-2]
    return text


def read_employee_excel(source) -> list:
    """
    读取员工信息 Excel（文件路径或文件流），自动识别中英文表头。
    所有列按文本处理以保留前导零；缺失的字段补空字符串。

    返回: [{'_row': Excel 行号, 'real_name': ..., 'id_number': ..., ...}, ...]
    异常: ValueError - 文件无法解析或找不到"姓名"列时抛出（消息为友好的中文提示）

    注意: 身份证号/银行账号等长数字请在 Excel 中以"文本"格式保存；
    数值格式的单元格超过 15 位会被 Excel 自身截断，任何工具都无法还原。
    """
    try:
        df = pd.read_excel(source, dtype=str).fillna('')
    except Exception:
        raise ValueError('文件解析失败，请确认是有效的 .xlsx 文件') from None

    # 表头标准化：别名 -> 标准字段名
    renamed = {}
    for col in df.columns:
        key = str(col).strip()
        if key in HEADER_ALIASES:
            renamed[col] = HEADER_ALIASES[key]
    df = df.rename(columns=renamed)

    if 'real_name' not in df.columns:
        raise ValueError('缺少"姓名"列（表头支持：姓名 / real_name）')

    rows = []
    for excel_row, record in enumerate(df.to_dict('records'), start=2):  # 数据从第 2 行起
        row = {'_row': excel_row}
        for field in IMPORT_COLUMNS:
            row[field] = _clean_cell(record.get(field, ''))
        rows.append(row)
    return rows


def classify_import_rows(rows: list, existing: dict = None) -> list:
    """
    结合数据库现状为每行数据打分类标签并生成校验信息。

    参数:
        rows:     read_employee_excel 的输出
        existing: {姓名: deleted} 映射，None 时自动从数据库查询
    返回:
        原列表，每行附加:
            status   - new(新增) / update(更新) / restore(恢复已删除) / error(错误，不导入)
            messages - 校验提示列表（错误与警告）
    """
    if existing is None:
        existing = {r['real_name']: r['deleted']
                    for r in query_all("SELECT real_name, deleted FROM employees")}

    seen = set()
    for row in rows:
        name = row['real_name']
        messages = []

        # 警告级校验（提示但不阻止导入）
        if row['id_number'] and len(row['id_number']) not in (15, 18):
            messages.append('身份证号位数异常（应为 15 或 18 位）')
        if row['phone'] and not row['phone'].isdigit():
            messages.append('电话包含非数字字符')

        # 分类（错误级会阻止该行导入）
        if not name:
            row['status'] = 'error'
            messages.insert(0, '姓名为空')
        elif name in seen:
            row['status'] = 'error'
            messages.insert(0, '文件中姓名重复')
        elif name not in existing:
            row['status'] = 'new'
        elif existing[name]:
            row['status'] = 'restore'
        else:
            row['status'] = 'update'

        seen.add(name)
        row['messages'] = messages
    return rows


def summarize_import_rows(rows: list) -> dict:
    """统计分类结果：{total, new, update, restore, error}。"""
    summary = {'total': len(rows), 'new': 0, 'update': 0, 'restore': 0, 'error': 0}
    for row in rows:
        summary[row['status']] += 1
    return summary


def import_employee_rows(rows: list) -> dict:
    """
    将分类后的行写入数据库（跳过 error/空姓名行，单事务）：
    新姓名插入；已存在更新；已软删除的同名记录自动恢复并清除审计信息。

    返回: {inserted, updated, restored, skipped}
    """
    existing = {r['real_name']: r['deleted']
                for r in query_all("SELECT real_name, deleted FROM employees")}

    inserted = updated = restored = skipped = 0
    with get_connection() as conn:
        for row in rows:
            if row.get('status') == 'error' or not row['real_name']:
                skipped += 1
                continue

            name = row['real_name']
            values = tuple(row[field] for field in IMPORT_COLUMNS[1:])
            sx = get_pinyin_sx(name)

            if name in existing:
                conn.execute('''
                    UPDATE employees
                    SET id_number=?, bank_account=?, bank_address=?, phone=?, sx=?,
                        deleted=0, deleted_at=NULL, deleted_by=NULL
                    WHERE real_name=?
                ''', (*values, sx, name))
                if existing[name]:
                    restored += 1
                else:
                    updated += 1
            else:
                conn.execute('''
                    INSERT INTO employees (real_name, id_number, bank_account, bank_address, phone, sx)
                    VALUES (?, ?, ?, ?, ?, ?)
                ''', (name, *values, sx))
                existing[name] = 0
                inserted += 1
        conn.commit()

    return {'inserted': inserted, 'updated': updated,
            'restored': restored, 'skipped': skipped}
