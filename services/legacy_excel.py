"""将旧版 BIFF/CFB .xls 员工表转换为供导入解析的内存工作簿。"""

import io
import struct

import xlrd
from xlrd.compdoc import CompDoc
from openpyxl import Workbook


OLE_MAGIC = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
MAX_LEGACY_CELLS = 1_000_000
FORMULA_PLACEHOLDER = '=XLS_FORMULA_UNSUPPORTED'


def _formula_positions(raw: bytes) -> set:
    """从 BIFF FORMULA 记录取得坐标；xlrd 的 cell.value 只有缓存值。"""
    container = CompDoc(raw)
    stream = container.get_named_stream('Workbook') or container.get_named_stream('Book')
    if not stream:
        raise ValueError('未找到旧版 Excel 工作簿数据')
    found = set()
    current_sheet = -1
    sheet_index = -1
    position = 0
    while position + 4 <= len(stream):
        record_type, length = struct.unpack_from('<HH', stream, position)
        if record_type == 0 or position + 4 + length > len(stream):
            break
        payload = position + 4
        if record_type in {0x0009, 0x0209, 0x0409, 0x0809} and length >= 4:
            substream_type = struct.unpack_from('<H', stream, payload + 2)[0]
            current_sheet = -1
            if substream_type == 0x0010:
                sheet_index += 1
                current_sheet = sheet_index
        elif record_type == 0x000A:
            current_sheet = -1
        elif record_type == 0x0006 and current_sheet >= 0 and length >= 4:
            row, column = struct.unpack_from('<HH', stream, payload)
            found.add((current_sheet, row, column))
        position = payload + length
    return found


def load_legacy_xls(raw: bytes) -> Workbook:
    """只用于数据导入；公式标记为不可导入，不尝试保留旧模板样式。"""
    if not raw.startswith(OLE_MAGIC):
        raise ValueError('文件不是有效的旧版 .xls 工作簿')
    try:
        formula_cells = _formula_positions(raw)
        source = xlrd.open_workbook(file_contents=raw, on_demand=True)
    except Exception as exc:
        raise ValueError('旧版 .xls 文件解析失败或已加密') from exc
    target = Workbook()
    target.remove(target.active)
    total_cells = 0
    try:
        if not source.nsheets:
            raise ValueError('旧版 .xls 文件没有工作表')
        for sheet_index in range(source.nsheets):
            old_sheet = source.sheet_by_index(sheet_index)
            total_cells += old_sheet.nrows * old_sheet.ncols
            if total_cells > MAX_LEGACY_CELLS:
                raise ValueError('旧版 .xls 数据单元格过多，无法安全导入')
            sheet = target.create_sheet(old_sheet.name)
            sheet.sheet_state = 'hidden' if old_sheet.visibility else 'visible'
            for row_index in range(old_sheet.nrows):
                for column_index in range(old_sheet.row_len(row_index)):
                    old_cell = old_sheet.cell(row_index, column_index)
                    if old_cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    cell = sheet.cell(row_index + 1, column_index + 1)
                    if (sheet_index, row_index, column_index) in formula_cells:
                        cell.value = FORMULA_PLACEHOLDER
                        continue
                    if old_cell.ctype == xlrd.XL_CELL_ERROR:
                        cell.value = FORMULA_PLACEHOLDER
                    elif old_cell.ctype == xlrd.XL_CELL_DATE:
                        cell.value = xlrd.xldate.xldate_as_datetime(old_cell.value, source.datemode)
                    elif old_cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        cell.value = bool(old_cell.value)
                    elif old_cell.ctype == xlrd.XL_CELL_NUMBER:
                        value = old_cell.value
                        cell.value = int(value) if value.is_integer() else value
                    else:
                        cell.value = old_cell.value
                        cell.data_type = 's'
        if all(sheet.sheet_state != 'visible' for sheet in target.worksheets):
            target.worksheets[0].sheet_state = 'visible'
        return target
    except Exception:
        target.close()
        raise
    finally:
        source.release_resources()


def load_import_workbook(source, *, data_only=False):
    """按实际文件签名解析 .xls 或 .xlsx，供导入入口复用。"""
    if isinstance(source, (bytes, bytearray)):
        raw = bytes(source)
    elif hasattr(source, 'read'):
        raw = source.read()
    else:
        with open(source, 'rb') as file:
            raw = file.read()
    if raw.startswith(OLE_MAGIC):
        return load_legacy_xls(raw)
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(raw), read_only=True, data_only=data_only)
