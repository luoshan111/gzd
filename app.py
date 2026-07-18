"""
员工信息管理系统 - Flask 主应用
功能：员工信息的增删改查、Excel导出、用户登录认证
"""

from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, session, flash
import sqlite3
import pandas as pd
from pypinyin import lazy_pinyin, Style
import io
import os
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash

# 创建Flask应用实例
app = Flask(__name__)
# 设置密钥，用于session加密
app.secret_key = 'gongzidan_secret_key_2026'

# 常量配置
DB_PATH = 'test.db'              # SQLite数据库文件路径
INPUT_XLSX = 'input.xlsx'        # 输入Excel文件路径
OUTPUT_XLSX = 'output.xlsx'      # 导出Excel文件路径
MANUAL_OUTPUT_XLSX = 'shuchu.xlsx'  # 手动导出Excel文件路径


def get_db_connection():
    """
    获取数据库连接
    返回: sqlite3.Connection对象
    """
    conn = sqlite3.connect(DB_PATH)
    # 设置行工厂，使查询结果可以通过列名访问
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    """
    初始化数据库
    创建employees表（员工信息）和users表（登录用户）
    """
    conn = sqlite3.connect(DB_PATH)
    # 创建员工信息表
    conn.execute('''
        CREATE TABLE IF NOT EXISTS employees(
            real_name TEXT PRIMARY KEY,    -- 姓名，主键
            id_number TEXT,                 -- 身份证号
            bank_account TEXT,              -- 银行账号
            bank_address TEXT,              -- 银行地址
            phone TEXT,                     -- 电话号码
            sx TEXT                         -- 拼音首字母，用于搜索
        )
    ''')
    # 创建用户表，用于登录认证
    conn.execute('''
        CREATE TABLE IF NOT EXISTS users(
            id INTEGER PRIMARY KEY AUTOINCREMENT,  -- 用户ID，自增
            username TEXT UNIQUE NOT NULL,          -- 用户名，唯一
            password_hash TEXT NOT NULL,            -- 密码哈希值
            is_admin INTEGER DEFAULT 0,             -- 是否管理员，0否1是
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP  -- 创建时间
        )
    ''')
    conn.commit()
    conn.close()


def get_pinyin_sx(name):
    """
    获取姓名的拼音首字母
    参数: name - 中文姓名
    返回: 拼音首字母字符串，如"张三"返回"zs"
    """
    result = lazy_pinyin(name, style=Style.FIRST_LETTER)
    return ''.join(char.lower() for char in result)


def login_required(f):
    """
    登录验证装饰器
    用于保护需要登录才能访问的路由
    未登录时：API请求返回JSON错误，页面请求重定向到登录页
    """
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            # 判断是API请求还是页面请求
            if request.is_json or request.path.startswith('/api/'):
                return jsonify({'success': False, 'message': '请先登录', 'need_login': True})
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function


def admin_required(f):
    """
    管理员权限验证装饰器
    用于保护只有管理员才能访问的路由
    """
    @wraps(f)
    def decorated_function(*args, **kwargs):
        # 先检查是否登录
        if 'user_id' not in session:
            if request.is_json or request.path.startswith('/api/'):
                return jsonify({'success': False, 'message': '请先登录', 'need_login': True})
            return redirect(url_for('login'))
        # 再检查是否是管理员
        if not session.get('is_admin'):
            return jsonify({'success': False, 'message': '需要管理员权限'}), 403
        return f(*args, **kwargs)
    return decorated_function


# ==================== 登录相关路由 ====================

@app.route('/login')
def login():
    """
    登录页面路由
    如果已登录则直接跳转到主页
    """
    if 'user_id' in session:
        return redirect(url_for('index'))
    return render_template('login.html')


@app.route('/api/login', methods=['POST'])
def api_login():
    """
    登录API接口
    接收JSON格式的用户名和密码，验证后设置session
    """
    data = request.get_json()
    username = data.get('username', '').strip()
    password = data.get('password', '')

    # 验证输入
    if not username or not password:
        return jsonify({'success': False, 'message': '请输入用户名和密码'})

    # 查询用户
    conn = get_db_connection()
    user = conn.execute(
        "SELECT id, username, password_hash, is_admin FROM users WHERE username = ?",
        (username,)
    ).fetchone()
    conn.close()

    # 验证密码
    if user and check_password_hash(user['password_hash'], password):
        # 登录成功，设置session
        session['user_id'] = user['id']
        session['username'] = user['username']
        session['is_admin'] = user['is_admin']
        return jsonify({'success': True, 'message': '登录成功', 'is_admin': user['is_admin']})
    else:
        return jsonify({'success': False, 'message': '用户名或密码错误'})


@app.route('/api/logout', methods=['POST'])
def api_logout():
    """
    登出API接口
    清除session中的用户信息
    """
    session.clear()
    return jsonify({'success': True, 'message': '已退出登录'})


@app.route('/api/check-login')
def check_login():
    """
    检查登录状态API
    返回当前用户的登录信息和权限
    """
    if 'user_id' in session:
        return jsonify({
            'success': True,
            'logged_in': True,
            'username': session.get('username'),
            'is_admin': session.get('is_admin', False)
        })
    return jsonify({'success': True, 'logged_in': False})


# ==================== 主页路由 ====================

@app.route('/')
@login_required
def index():
    """
    主页路由
    需要登录才能访问
    """
    return render_template('index.html')


# ==================== 员工信息API ====================

@app.route('/api/search', methods=['GET'])
@login_required
def search():
    """
    搜索员工信息API
    通过姓名或拼音首字母模糊搜索
    """
    name = request.args.get('name', '').strip()
    if not name:
        return jsonify({'success': False, 'message': '请输入查询关键词'})

    conn = get_db_connection()
    # 模糊查询姓名和拼音首字母
    cursor = conn.execute(
        "SELECT * FROM employees WHERE real_name LIKE ? OR sx LIKE ?",
        (f'%{name}%', f'%{name}%')
    )
    results = [dict(row) for row in cursor.fetchall()]
    conn.close()

    if results:
        return jsonify({'success': True, 'data': results})
    else:
        return jsonify({'success': False, 'message': '未找到匹配的记录'})


@app.route('/api/employees', methods=['GET'])
@login_required
def get_employees():
    """
    获取所有员工列表API
    按姓名排序返回
    """
    conn = get_db_connection()
    cursor = conn.execute("SELECT * FROM employees ORDER BY real_name")
    results = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return jsonify({'success': True, 'data': results})


@app.route('/api/employee', methods=['POST', 'PUT'])
@login_required
def add_employee():
    """
    添加或更新员工信息API
    如果姓名已存在则更新，否则新增
    """
    data = request.get_json()

    real_name = data.get('real_name', '').strip()
    id_number = data.get('id_number', '').strip()
    bank_account = data.get('bank_account', '').strip()
    bank_address = data.get('bank_address', '').strip()
    phone = data.get('phone', '').strip()

    # 姓名为必填项
    if not real_name:
        return jsonify({'success': False, 'message': '姓名为必填项'})

    # 计算拼音首字母
    sx = get_pinyin_sx(real_name)

    conn = get_db_connection()
    # 检查员工是否已存在
    exists = conn.execute("SELECT 1 FROM employees WHERE real_name = ?", (real_name,)).fetchone()

    if exists:
        # 更新已有员工信息
        conn.execute('''
            UPDATE employees
            SET id_number=?, bank_account=?, bank_address=?, phone=?, sx=?
            WHERE real_name=?
        ''', (id_number, bank_account, bank_address, phone, sx, real_name))
        message = '信息更新成功'
    else:
        # 新增员工信息
        conn.execute('''
            INSERT INTO employees (real_name, id_number, bank_account, bank_address, phone, sx)
            VALUES (?, ?, ?, ?, ?, ?)
        ''', (real_name, id_number, bank_account, bank_address, phone, sx))
        message = '信息添加成功'

    conn.commit()
    conn.close()

    return jsonify({'success': True, 'message': message})


@app.route('/api/employee/<name>', methods=['DELETE'])
@login_required
def delete_employee(name):
    """
    删除员工信息API
    参数: name - 员工姓名（URL参数）
    """
    conn = get_db_connection()
    conn.execute("DELETE FROM employees WHERE real_name = ?", (name,))
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': '删除成功'})


# ==================== Excel导出API ====================

@app.route('/api/export', methods=['POST'])
@login_required
def export_excel():
    """
    从input.xlsx导出员工信息到output.xlsx
    读取input.xlsx中的姓名列表，匹配数据库中的信息后导出
    """
    try:
        # 检查输入文件是否存在
        if not os.path.exists(INPUT_XLSX):
            return jsonify({'success': False, 'message': 'input.xlsx 文件不存在'})

        import openpyxl
        wb = openpyxl.load_workbook(INPUT_XLSX)
        ws = wb.active

        # 读取所有姓名（第一列）
        names = []
        for row in ws.iter_rows(min_row=1, values_only=True):
            if row[0]:
                names.append(row[0])

        if not names:
            return jsonify({'success': False, 'message': 'input.xlsx 中没有姓名数据'})

        # 创建输出DataFrame
        data = {
            '姓名': names,
            '身份证号': [''] * len(names),
            '银行账号': [''] * len(names),
            '银行地址': [''] * len(names),
            '电话': [''] * len(names)
        }
        df = pd.DataFrame(data)

        # 从数据库填充信息
        conn = get_db_connection()
        sql_real_name = conn.execute("SELECT real_name FROM employees").fetchall()
        existing_names = {row[0] for row in sql_real_name}

        # 遍历姓名，匹配数据库中的信息
        for i in range(len(df)):
            name = df.iloc[i]['姓名']
            if name in existing_names:
                cursor = conn.execute(
                    "SELECT id_number, bank_account, bank_address, phone FROM employees WHERE real_name = ?",
                    (name,)
                )
                result = cursor.fetchone()
                if result:
                    df.at[i, '身份证号'] = result[0]
                    df.at[i, '银行账号'] = result[1]
                    df.at[i, '银行地址'] = result[2]
                    df.at[i, '电话'] = result[3]

        conn.close()
        # 保存到Excel文件
        df.to_excel(OUTPUT_XLSX, index=False, engine='openpyxl')

        return jsonify({'success': True, 'message': '导出成功', 'file': OUTPUT_XLSX})
    except Exception as e:
        return jsonify({'success': False, 'message': f'导出失败: {str(e)}'})


@app.route('/api/download/<filename>')
@login_required
def download_file(filename):
    """
    文件下载API
    用于下载导出的Excel文件
    """
    return send_file(filename, as_attachment=True)


@app.route('/api/export-manual', methods=['POST'])
@login_required
def export_manual_excel():
    """
    手动输入姓名导出Excel
    接收姓名列表，从数据库匹配信息后导出到shuchu.xlsx
    """
    try:
        data = request.get_json()
        names = data.get('names', [])

        if not names:
            return jsonify({'success': False, 'message': '请提供姓名列表'})

        result_data = []
        conn = get_db_connection()

        # 遍历姓名，查询数据库
        for name in names:
            cursor = conn.execute(
                "SELECT id_number, bank_account, bank_address, phone FROM employees WHERE real_name = ?",
                (name,)
            )
            result = cursor.fetchone()

            if result:
                # 找到匹配的员工信息
                result_data.append({
                    '姓名': name,
                    '身份证号': result[0] or '',
                    '银行账号': result[1] or '',
                    '银行地址': result[2] or '',
                    '电话': result[3] or ''
                })
            else:
                # 数据库中无此姓名，导出空信息
                result_data.append({
                    '姓名': name,
                    '身份证号': '',
                    '银行账号': '',
                    '银行地址': '',
                    '电话': ''
                })

        conn.close()
        # 创建DataFrame并保存
        df = pd.DataFrame(result_data)
        df.to_excel(MANUAL_OUTPUT_XLSX, index=False, engine='openpyxl')

        return jsonify({'success': True, 'message': '导出成功', 'file': MANUAL_OUTPUT_XLSX})
    except Exception as e:
        return jsonify({'success': False, 'message': f'导出失败: {str(e)}'})


@app.route('/api/input-preview')
@login_required
def input_preview():
    """
    预览input.xlsx内容API
    返回前20行姓名数据，用于前端展示
    """
    try:
        if os.path.exists(INPUT_XLSX):
            import openpyxl
            wb = openpyxl.load_workbook(INPUT_XLSX)
            ws = wb.active

            # 读取前20行数据
            names = []
            for row in ws.iter_rows(max_row=20, values_only=True):
                if row[0]:
                    names.append({'姓名': row[0]})

            columns = ['姓名']
            return jsonify({'success': True, 'data': names, 'columns': columns})
        else:
            return jsonify({'success': False, 'message': 'input.xlsx 不存在'})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)})


# ==================== 管理员功能API ====================

@app.route('/admin')
@admin_required
def admin_page():
    """
    管理员后台页面
    只有管理员才能访问
    """
    return render_template('admin.html')


@app.route('/api/admin/users', methods=['GET'])
@admin_required
def get_users():
    """
    获取所有用户列表API
    仅管理员可访问
    """
    conn = get_db_connection()
    users = conn.execute("SELECT id, username, is_admin, created_at FROM users ORDER BY id").fetchall()
    conn.close()
    return jsonify({'success': True, 'data': [dict(u) for u in users]})


@app.route('/api/admin/user', methods=['POST'])
@admin_required
def create_user():
    """
    创建新用户API
    仅管理员可访问
    """
    data = request.get_json()
    username = data.get('username', '').strip()
    password = data.get('password', '')
    is_admin = data.get('is_admin', 0)

    # 验证输入
    if not username or not password:
        return jsonify({'success': False, 'message': '用户名和密码不能为空'})

    if len(password) < 4:
        return jsonify({'success': False, 'message': '密码至少4位'})

    # 对密码进行哈希处理
    password_hash = generate_password_hash(password)

    conn = get_db_connection()
    try:
        conn.execute(
            "INSERT INTO users (username, password_hash, is_admin) VALUES (?, ?, ?)",
            (username, password_hash, is_admin)
        )
        conn.commit()
        conn.close()
        return jsonify({'success': True, 'message': '用户创建成功'})
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({'success': False, 'message': '用户名已存在'})


@app.route('/api/admin/user/<int:user_id>', methods=['DELETE'])
@admin_required
def delete_user(user_id):
    """
    删除用户API
    仅管理员可访问，不能删除自己的账户
    """
    # 防止删除自己的账户
    if user_id == session.get('user_id'):
        return jsonify({'success': False, 'message': '不能删除自己的账户'})

    conn = get_db_connection()
    conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': '用户删除成功'})


@app.route('/api/admin/user/<int:user_id>/password', methods=['PUT'])
@admin_required
def reset_user_password(user_id):
    """
    重置用户密码API
    仅管理员可访问
    """
    data = request.get_json()
    password = data.get('password', '')

    # 验证密码长度
    if not password or len(password) < 4:
        return jsonify({'success': False, 'message': '密码至少4位'})

    # 对新密码进行哈希处理
    password_hash = generate_password_hash(password)

    conn = get_db_connection()
    conn.execute("UPDATE users SET password_hash = ? WHERE id = ?", (password_hash, user_id))
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'message': '密码重置成功'})


# ==================== 应用启动 ====================

if __name__ == '__main__':
    # 初始化数据库
    init_db()
    # 启动Flask开发服务器
    # debug=True 开启调试模式
    # host='0.0.0.0' 允许外部访问
    # port=5001 使用5001端口
    app.run(debug=True, host='0.0.0.0', port=5001)
