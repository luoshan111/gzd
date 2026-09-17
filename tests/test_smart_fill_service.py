import hashlib
import io
import json
import unittest
from unittest.mock import Mock, patch

from flask import Flask
from openpyxl import Workbook, load_workbook

import import_export_routes
import smart_fill_service as service


def workbook_bytes(setup):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '工资表'
    setup(worksheet)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def analysis(columns, data_end_row=None):
    return {
        'confidence': 0.94,
        'tables': [{
            'sheet': '工资表',
            'header_row': 2,
            'data_start_row': 3,
            'data_end_row': data_end_row,
            'columns': columns,
        }],
        'notes': [],
    }


class SmartFillServiceTests(unittest.TestCase):
    def test_ai_snapshot_contains_headers_but_no_template_data_or_sheet_name(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '张三项目工资表'
        worksheet.append(['姓名', '身份证号', '应发工资'])
        worksheet.append(['张三', '110101199001011234', 8888])

        snapshot_text = json.dumps(
            service._build_template_snapshot(workbook), ensure_ascii=False
        )

        self.assertIn('sheet_1', snapshot_text)
        self.assertIn('姓名', snapshot_text)
        self.assertIn('身份证号', snapshot_text)
        self.assertNotIn('张三项目工资表', snapshot_text)
        self.assertNotIn('张三', snapshot_text)
        self.assertNotIn('110101199001011234', snapshot_text)
        self.assertNotIn('8888', snapshot_text)
        workbook.close()

    def test_anonymous_sheet_id_is_resolved_only_locally(self):
        workbook = Workbook()
        workbook.active.title = '内部项目表'
        raw = {
            'confidence': 0.9,
            'tables': [{'sheet_id': 'sheet_1', 'header_row': 1}],
        }

        resolved = service._resolve_sheet_ids(workbook, raw)

        self.assertEqual(resolved['tables'][0]['sheet'], '内部项目表')
        workbook.close()

    def test_mapping_to_unrelated_header_is_rejected(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '工资表'
        worksheet.append(['姓名', '部门'])
        raw = {
            'confidence': 0.99,
            'tables': [{
                'sheet': '工资表', 'header_row': 1, 'data_start_row': 2,
                'data_end_row': None,
                'columns': {'real_name': 'A', 'id_number': 'B'},
            }],
        }

        with self.assertRaisesRegex(service.TemplateAnalysisError, '未可靠识别'):
            service._validate_analysis(workbook, raw)
        workbook.close()

    def test_mapping_to_blank_far_column_is_rejected(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '工资表'
        worksheet.append(['姓名', '身份证号'])
        raw = {
            'confidence': 0.99,
            'tables': [{
                'sheet': '工资表', 'header_row': 1, 'data_start_row': 2,
                'data_end_row': None,
                'columns': {'real_name': 'A', 'id_number': 'XFD'},
            }],
        }

        with self.assertRaisesRegex(service.TemplateAnalysisError, '未可靠识别'):
            service._validate_analysis(workbook, raw)
        workbook.close()

    def test_validated_mapping_keeps_exact_header_evidence_for_preview(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '工资表'
        worksheet.append(['人员', '证件信息', '结算账户', '收款银行', '联络方式'])
        raw = {
            'confidence': 0.91,
            'tables': [{
                'sheet': '工资表', 'header_row': 1, 'data_start_row': 2,
                'data_end_row': None,
                'columns': {
                    'real_name': 'A', 'id_number': 'B', 'bank_account': 'C',
                    'bank_address': 'D', 'phone': 'E',
                },
            }],
        }

        validated = service._validate_analysis(workbook, raw)

        headers = validated['tables'][0]['target_headers']
        self.assertEqual(headers['real_name'], {'row': 1, 'text': '人员'})
        self.assertEqual(headers['bank_account'], {'row': 1, 'text': '结算账户'})
        workbook.close()

    def test_local_fallback_recognizes_common_payroll_headers(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '工资明细'
        worksheet.append(['序号', '员工姓名', '身份证号码', '工资卡号', '开户银行', '联系电话'])

        result = service._heuristic_analysis(workbook)

        self.assertEqual(result['tables'][0]['sheet'], '工资明细')
        self.assertEqual(result['tables'][0]['data_start_row'], 2)
        self.assertEqual(result['tables'][0]['columns']['real_name'], 'B')
        self.assertEqual(result['tables'][0]['columns']['bank_account'], 'D')
        workbook.close()

    def test_invalid_ai_mapping_falls_back_to_strict_local_headers(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '工资明细'
        worksheet.append(['姓名', '身份证号'])
        response = Mock()
        response.choices = [Mock(message=Mock(content='不是 JSON'))]
        client = Mock()
        client.chat.completions.create.return_value = response

        with patch.object(service.config, 'AGENT_API_KEY', 'test-key'), \
             patch.object(service, 'get_client', return_value=client):
            result = service.analyze_template(workbook)

        self.assertEqual(result['confidence'], 0.82)
        self.assertEqual(result['tables'][0]['columns']['real_name'], 'A')
        workbook.close()

    def test_rejects_xlsx_with_excessive_uncompressed_size(self):
        template = workbook_bytes(lambda ws: ws.append(['姓名', '身份证号']))

        with patch.object(service, 'MAX_XLSX_UNCOMPRESSED_BYTES', 1):
            with self.assertRaisesRegex(service.TemplateFillError, '解压后体积过大'):
                service.prepare_template(template)

    def test_fills_existing_name_rows_and_preserves_formula(self):
        template = workbook_bytes(lambda ws: (
            ws.append(['项目工资表']),
            ws.append(['序号', '姓名', '身份证号', '银行卡号', '备注']),
            ws.append([1, '张三', None, None, '=1+1']),
        ))
        employees = {
            '张三': {
                'real_name': '张三', 'id_number': '110101199001011234',
                'bank_account': '001234567890', 'bank_address': '测试银行',
                'phone': '13800000000',
            }
        }
        plan = analysis({'real_name': 'B', 'id_number': 'C', 'bank_account': 'D'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value=employees):
            result, summary, _ = service.fill_template(template)

        workbook = load_workbook(io.BytesIO(result), data_only=False)
        worksheet = workbook['工资表']
        self.assertEqual(worksheet['B3'].value, '张三')
        self.assertEqual(worksheet['C3'].value, '110101199001011234')
        self.assertEqual(worksheet['D3'].value, '001234567890')
        self.assertEqual(worksheet['E3'].value, '=1+1')
        self.assertEqual(worksheet['C3'].number_format, '@')
        self.assertEqual(summary['matched'], 1)
        self.assertEqual(summary['cells_filled'], 2)
        workbook.close()

    def test_uses_supplied_names_and_reports_database_misses(self):
        template = workbook_bytes(lambda ws: (
            ws.append(['项目工资表']),
            ws.append(['姓名', '身份证号', '开户行']),
            ws.append([None, None, None]),
        ))
        employees = {
            '张三': {
                'real_name': '张三', 'id_number': '110101199001011234',
                'bank_account': '', 'bank_address': '测试银行', 'phone': '',
            }
        }
        plan = analysis({'real_name': 'A', 'id_number': 'B', 'bank_address': 'C'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value=employees):
            result, summary, _ = service.fill_template(template, names=['张三', '不存在'])

        workbook = load_workbook(io.BytesIO(result))
        worksheet = workbook['工资表']
        self.assertEqual(worksheet['A3'].value, '张三')
        self.assertEqual(worksheet['B3'].value, '110101199001011234')
        self.assertEqual(worksheet['C3'].value, '测试银行')
        self.assertEqual(summary['matched'], 1)
        self.assertEqual(summary['missing'], ['不存在'])
        workbook.close()

    def test_supplied_name_does_not_claim_a_partially_populated_row(self):
        def setup(worksheet):
            worksheet.append(['项目工资表'])
            worksheet.append(['姓名', '身份证号', '开户行'])
            worksheet.append([None, '已有身份证', None])
            worksheet['A4'].number_format = '@'

        template = workbook_bytes(setup)
        employees = {
            '张三': {
                'real_name': '张三', 'id_number': '数据库身份证',
                'bank_account': '', 'bank_address': '测试银行', 'phone': '',
            }
        }
        plan = analysis({'real_name': 'A', 'id_number': 'B', 'bank_address': 'C'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value=employees):
            result, _, _ = service.fill_template(template, names=['张三'])

        workbook = load_workbook(io.BytesIO(result))
        worksheet = workbook['工资表']
        self.assertIsNone(worksheet['A3'].value)
        self.assertEqual(worksheet['B3'].value, '已有身份证')
        self.assertEqual(worksheet['A4'].value, '张三')
        self.assertEqual(worksheet['B4'].value, '数据库身份证')
        workbook.close()

    def test_supplied_names_stop_before_summary_row(self):
        template = workbook_bytes(lambda ws: (
            ws.append(['项目工资表']),
            ws.append(['姓名', '身份证号', '备注']),
            ws.append([None, None, None]),
            ws.append([None, None, '合计']),
        ))
        employees = {
            name: {
                'real_name': name, 'id_number': f'{name}-身份证',
                'bank_account': '', 'bank_address': '', 'phone': '',
            }
            for name in ('张三', '李四')
        }
        plan = analysis({'real_name': 'A', 'id_number': 'B'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value=employees):
            with self.assertRaisesRegex(service.TemplateFillError, '空白行不足'):
                service.fill_template(template, names=['张三', '李四'])

    def test_appended_rows_copy_translated_formulas_and_request_recalculation(self):
        def setup(worksheet):
            worksheet.append(['项目工资表'])
            worksheet.append(['姓名', '身份证号', '日工资', '总工资'])
            worksheet.append([None, None, 300, '=C3*2'])
            worksheet['D3'].number_format = '0.00'

        template = workbook_bytes(setup)
        employees = {
            name: {
                'real_name': name, 'id_number': f'{name}-身份证',
                'bank_account': '', 'bank_address': '', 'phone': '',
            }
            for name in ('张三', '李四')
        }
        plan = analysis({'real_name': 'A', 'id_number': 'B'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value=employees):
            result, _, _ = service.fill_template(template, names=['张三', '李四'])

        workbook = load_workbook(io.BytesIO(result), data_only=False)
        worksheet = workbook['工资表']
        self.assertEqual(worksheet['D3'].value, '=C3*2')
        self.assertEqual(worksheet['D4'].value, '=C4*2')
        self.assertIsNone(worksheet['C4'].value)
        self.assertEqual(worksheet['D4'].number_format, '0.00')
        self.assertEqual(workbook.calculation.calcMode, 'auto')
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        self.assertTrue(workbook.calculation.forceFullCalc)
        self.assertTrue(workbook.calculation.calcOnSave)
        workbook.close()

    def test_supplied_name_updates_every_existing_duplicate_row(self):
        template = workbook_bytes(lambda ws: (
            ws.append(['项目工资表']),
            ws.append(['姓名', '身份证号']),
            ws.append(['张三', None]),
            ws.append(['张三', None]),
        ))
        employees = {
            '张三': {
                'real_name': '张三', 'id_number': '数据库身份证',
                'bank_account': '', 'bank_address': '', 'phone': '',
            }
        }
        plan = analysis({'real_name': 'A', 'id_number': 'B'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value=employees):
            result, _, _ = service.fill_template(template, names=['张三'])

        workbook = load_workbook(io.BytesIO(result))
        worksheet = workbook['工资表']
        self.assertEqual(worksheet['B3'].value, '数据库身份证')
        self.assertEqual(worksheet['B4'].value, '数据库身份证')
        workbook.close()

    def test_does_not_overwrite_existing_target_value(self):
        template = workbook_bytes(lambda ws: (
            ws.append(['项目工资表']),
            ws.append(['姓名', '身份证号', '开户行']),
            ws.append(['张三', '人工确认值', None]),
        ))
        employees = {
            '张三': {
                'real_name': '张三', 'id_number': '数据库值',
                'bank_account': '', 'bank_address': '测试银行', 'phone': '',
            }
        }
        plan = analysis({'real_name': 'A', 'id_number': 'B', 'bank_address': 'C'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value=employees):
            result, summary, _ = service.fill_template(template)

        workbook = load_workbook(io.BytesIO(result))
        worksheet = workbook['工资表']
        self.assertEqual(worksheet['B3'].value, '人工确认值')
        self.assertEqual(worksheet['C3'].value, '测试银行')
        self.assertEqual(summary['skipped_cells'], 1)
        workbook.close()

    def test_database_text_cannot_become_an_excel_formula(self):
        template = workbook_bytes(lambda ws: (
            ws.append(['项目工资表']),
            ws.append(['姓名', '身份证号', '开户行']),
            ws.append([None, None, None]),
        ))
        employees = {
            '=危险姓名': {
                'real_name': '=危险姓名', 'id_number': '=1+1',
                'bank_account': '', 'bank_address': '=HYPERLINK("bad")', 'phone': '',
            }
        }
        plan = analysis({'real_name': 'A', 'id_number': 'B', 'bank_address': 'C'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value=employees):
            result, _, _ = service.fill_template(template, names=['=危险姓名'])

        workbook = load_workbook(io.BytesIO(result), data_only=False)
        worksheet = workbook['工资表']
        for coordinate in ('A3', 'B3', 'C3'):
            self.assertEqual(worksheet[coordinate].data_type, 's')
        self.assertEqual(worksheet['B3'].value, '=1+1')
        self.assertEqual(worksheet['C3'].value, '=HYPERLINK("bad")')
        workbook.close()

    def test_stops_when_template_has_no_names_and_no_name_list(self):
        template = workbook_bytes(lambda ws: (
            ws.append(['项目工资表']),
            ws.append(['姓名', '身份证号']),
            ws.append([None, None]),
        ))
        plan = analysis({'real_name': 'A', 'id_number': 'B'})

        with patch.object(service, 'analyze_template', return_value=plan), \
             patch.object(service, '_employee_records', return_value={}):
            with self.assertRaisesRegex(service.TemplateFillError, '姓名列'):
                service.fill_template(template)


class SmartFillRouteTests(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.secret_key = 'test-secret'
        import_export_routes.register_import_export_routes(self.app)
        self.client = self.app.test_client()
        with self.client.session_transaction() as session:
            session['user_id'] = 1
            session['username'] = 'tester'
        import_export_routes._smart_fill_preview_store.clear()
        import_export_routes._smart_import_preview_store.clear()

    def test_smart_fill_preview_then_confirm_returns_download(self):
        analysis_result = {
            'confidence': 0.91,
            'tables': [{'sheet': '工资表', 'columns': {'real_name': 'A', 'id_number': 'B'}}],
        }
        preview_result = {
            'confidence': 0.91,
            'tables': [{
                'sheet': '工资表', 'header_row': 1, 'data_start_row': 2,
                'data_end_row': None,
                'fields': [
                    {'field': 'real_name', 'label': '姓名', 'column': 'A', 'header': '姓名', 'header_row': 1},
                    {'field': 'id_number', 'label': '身份证号', 'column': 'B', 'header': '身份证号', 'header_row': 1},
                ],
            }],
            'notes': [],
        }
        fill_result = (
            b'filled-workbook',
            {'matched': 1, 'missing': [], 'cells_filled': 3},
            {'confidence': 0.91},
        )
        with patch.object(import_export_routes, 'prepare_template',
                          return_value=(analysis_result, preview_result)):
            preview_response = self.client.post('/api/smart-fill-preview', data={
                'file': (io.BytesIO(b'template'), 'template.xlsx'),
                'names': json.dumps(['张三'], ensure_ascii=False),
            }, content_type='multipart/form-data')

        self.assertEqual(preview_response.status_code, 200)
        preview_json = preview_response.get_json()
        self.assertTrue(preview_json['preview_token'])
        self.assertEqual(preview_json['preview']['tables'][0]['fields'][1]['column'], 'B')
        stored_preview = import_export_routes._smart_fill_preview_store[
            preview_json['preview_token']
        ]
        self.assertNotIn('template_bytes', stored_preview)
        self.assertEqual(
            stored_preview['template_digest'], hashlib.sha256(b'template').hexdigest(),
        )

        with patch.object(import_export_routes, 'fill_template', return_value=fill_result) as fill_mock:
            response = self.client.post('/api/smart-fill-confirm', data={
                'preview_token': preview_json['preview_token'],
                'file': (io.BytesIO(b'template'), 'template.xlsx'),
                'names': json.dumps(['篡改后的姓名'], ensure_ascii=False),
                'analysis': json.dumps({'columns': {'id_number': 'XFD'}}),
            }, content_type='multipart/form-data')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, b'filled-workbook')
        self.assertEqual(response.headers['X-Smart-Fill-Matched'], '1')
        self.assertEqual(response.headers['X-Smart-Fill-Cells'], '3')
        fill_mock.assert_called_once_with(
            b'template', names=['张三'], analysis=analysis_result,
        )
        self.assertNotIn(preview_json['preview_token'],
                         import_export_routes._smart_fill_preview_store)

    def test_route_end_to_end_uses_the_client_template(self):
        template = workbook_bytes(lambda ws: (
            ws.append(['项目工资表']),
            ws.append(['姓名', '身份证号', '开户行']),
            ws.append([None, None, None]),
        ))
        employees = {
            '张三': {
                'real_name': '张三', 'id_number': '数据库身份证',
                'bank_account': '', 'bank_address': '测试银行', 'phone': '',
            }
        }
        with patch.object(service.config, 'AGENT_API_KEY', ''), \
             patch.object(service, '_employee_records', return_value=employees):
            preview_response = self.client.post('/api/smart-fill-preview', data={
                'file': (io.BytesIO(template), '客户端模板.xlsx'),
                'names': json.dumps(['张三'], ensure_ascii=False),
            }, content_type='multipart/form-data')
            self.assertEqual(preview_response.status_code, 200)
            token = preview_response.get_json()['preview_token']

            confirm_response = self.client.post('/api/smart-fill-confirm', data={
                'preview_token': token,
                'file': (io.BytesIO(template), '客户端模板.xlsx'),
            }, content_type='multipart/form-data')

        self.assertEqual(confirm_response.status_code, 200)
        workbook = load_workbook(io.BytesIO(confirm_response.data), data_only=False)
        worksheet = workbook['工资表']
        self.assertEqual(worksheet['A3'].value, '张三')
        self.assertEqual(worksheet['B3'].value, '数据库身份证')
        self.assertEqual(worksheet['C3'].value, '测试银行')
        workbook.close()

    def test_direct_smart_fill_cannot_bypass_preview(self):
        response = self.client.post('/api/smart-fill')

        self.assertEqual(response.status_code, 409)

    def test_preview_token_is_bound_to_current_user(self):
        token = import_export_routes._store_smart_fill_preview(
            1, hashlib.sha256(b'template').hexdigest(), [], {'confidence': 0.9},
        )
        with self.client.session_transaction() as session:
            session['user_id'] = 2
            session['username'] = 'other-user'

        response = self.client.post('/api/smart-fill-confirm', data={
            'preview_token': token,
            'file': (io.BytesIO(b'template'), 'template.xlsx'),
        }, content_type='multipart/form-data')

        self.assertEqual(response.status_code, 404)

    def test_confirm_rejects_a_different_client_template(self):
        token = import_export_routes._store_smart_fill_preview(
            1, hashlib.sha256(b'previewed-template').hexdigest(), [], {'confidence': 0.9},
        )
        with patch.object(import_export_routes, 'fill_template') as fill_mock:
            response = self.client.post('/api/smart-fill-confirm', data={
                'preview_token': token,
                'file': (io.BytesIO(b'different-template'), 'template.xlsx'),
            }, content_type='multipart/form-data')

        self.assertEqual(response.status_code, 409)
        fill_mock.assert_not_called()

    def test_preview_rejects_an_excessive_name_list(self):
        names = [f'姓名{i}' for i in range(import_export_routes.MAX_FILL_ROWS + 1)]
        with patch.object(import_export_routes, 'prepare_template') as prepare_mock:
            response = self.client.post('/api/smart-fill-preview', data={
                'file': (io.BytesIO(b'template'), 'template.xlsx'),
                'names': json.dumps(names, ensure_ascii=False),
            }, content_type='multipart/form-data')

        self.assertEqual(response.status_code, 400)
        prepare_mock.assert_not_called()


if __name__ == '__main__':
    unittest.main()
