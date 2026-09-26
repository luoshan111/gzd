import io
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from flask import Flask
from openpyxl import Workbook, load_workbook

from core import db
from routes.employee_routes import register_employee_routes
from services import employee_service, excel_utils, smart_fill_service, smart_import_service


class DatabaseAccessTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix='gzd-db-access-')
        self.addCleanup(directory.cleanup)
        self.database_path = Path(directory.name) / '员工数据 # test.db'
        override = patch.object(db.config, 'DB_PATH', str(self.database_path))
        override.start()
        self.addCleanup(override.stop)
        db.init_db()
        self.employee = {
            'real_name': '张三', 'id_number': '原身份证', 'bank_account': '000123',
            'bank_address': '原银行', 'phone': '13800000000',
        }
        employee_service.save_employee(self.employee, 'tester')

    def test_read_helpers_block_writes_and_database_remains_writable(self):
        for query in (db.query_all, db.query_one):
            with self.subTest(query=query.__name__):
                with self.assertRaises(sqlite3.OperationalError):
                    query("UPDATE employees SET phone='不允许' RETURNING *")
        self.assertEqual(db.query_one('SELECT phone FROM employees')['phone'], '13800000000')
        self.assertEqual(db.execute('UPDATE employees SET phone=? WHERE real_name=?',
                                    ('13900000000', '张三')), 1)
        self.assertEqual(db.query_one('SELECT phone FROM employees')['phone'], '13900000000')

    def test_read_only_connection_cannot_write_even_if_query_only_is_disabled(self):
        connection = db.get_connection(read_only=True)
        try:
            self.assertEqual(connection.execute('PRAGMA query_only').fetchone()[0], 1)
            connection.execute('PRAGMA query_only=OFF')
            with self.assertRaises(sqlite3.OperationalError):
                connection.execute('DELETE FROM employees')
        finally:
            connection.close()
        with self.assertRaises(ValueError):
            with db.db_session(commit=True, read_only=True):
                self.fail('只读会话不能提交')

    def test_read_only_connection_does_not_create_a_missing_database(self):
        missing = self.database_path.with_name('不存在.db')
        with patch.object(db.config, 'DB_PATH', str(missing)):
            with self.assertRaises(sqlite3.OperationalError):
                db.get_connection(read_only=True)
        self.assertFalse(missing.exists())

    def test_employee_edit_routes_still_create_update_delete_and_restore(self):
        app = Flask(__name__)
        app.secret_key = 'test-key'
        register_employee_routes(app)
        client = app.test_client()
        with client.session_transaction() as session:
            session['user_id'] = 1
            session['username'] = 'tester'
        employee = {**self.employee, 'real_name': '李四'}
        self.assertEqual(client.post('/api/employee', json=employee).status_code, 200)
        employee['phone'] = '13900000000'
        self.assertEqual(client.put('/api/employee', json=employee).status_code, 200)
        self.assertEqual(db.query_one('SELECT phone FROM employees WHERE real_name=?',
                                      ('李四',))['phone'], '13900000000')
        self.assertEqual(client.delete('/api/employee/李四').status_code, 200)
        self.assertEqual(db.query_one('SELECT deleted FROM employees WHERE real_name=?',
                                      ('李四',))['deleted'], 1)
        self.assertEqual(client.post('/api/employee/李四/restore').status_code, 200)
        self.assertEqual(db.query_one('SELECT deleted FROM employees WHERE real_name=?',
                                      ('李四',))['deleted'], 0)

    def test_export_and_template_fill_use_only_read_connections(self):
        workbook = Workbook()
        workbook.active.append(['姓名', '身份证号'])
        workbook.active.append(['模板旧人', '模板旧证件'])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        plan = {
            'confidence': 0.95, 'tables': [{
                'sheet': 'Sheet', 'header_row': 1, 'data_start_row': 2,
                'data_end_row': None, 'columns': {'real_name': 'A', 'id_number': 'B'},
            }],
        }
        before = db.query_all('SELECT * FROM employees')
        with patch.object(db, 'get_connection', wraps=db.get_connection) as connections:
            result, _, _ = smart_fill_service.fill_template(
                output.getvalue(), names=['张三'], analysis=plan, overwrite=True,
            )
            records = excel_utils.build_employee_records(['张三'])
            self.assertEqual(employee_service.list_employees(1, 20)['total'], 1)
            self.assertEqual(len(employee_service.search_employees('张三')), 1)
        self.assertGreaterEqual(connections.call_count, 4)
        for call in connections.call_args_list:
            self.assertEqual(call.kwargs, {'read_only': True})
        self.assertEqual(records[0]['姓名'], '张三')
        workbook = load_workbook(io.BytesIO(result))
        self.assertEqual(workbook.active['A2'].value, '张三')
        self.assertEqual(workbook.active['B2'].value, '原身份证')
        workbook.close()
        self.assertEqual(db.query_all('SELECT * FROM employees'), before)

    def test_standard_and_smart_import_still_insert_and_update(self):
        standard_rows = [
            {'real_name': '张三', 'id_number': '标准新身份证', 'bank_account': '000456',
             'bank_address': '新银行', 'phone': '13900000000', 'status': 'update'},
            {'real_name': '李四', 'id_number': '李四证件', 'bank_account': '',
             'bank_address': '', 'phone': '', 'status': 'new'},
        ]
        result = excel_utils.import_employee_rows(standard_rows)
        self.assertEqual((result['inserted'], result['updated']), (1, 1))
        workbook = Workbook()
        workbook.active.append(['姓名', '身份证号'])
        workbook.active.append(['张三', '智能新身份证'])
        workbook.active.append(['王五', '王五证件'])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        plan = {'confidence': 0.95, 'tables': [{
            'sheet': 'Sheet', 'header_row': 1, 'data_start_row': 2,
            'data_end_row': None, 'columns': {'real_name': 'A', 'id_number': 'B'},
        }]}
        result = smart_import_service.execute_smart_import(
            output.getvalue(), plan, {'new', 'update'},
        )
        self.assertEqual((result['inserted'], result['updated']), (1, 1))
        zhang = db.query_one('SELECT * FROM employees WHERE real_name=?', ('张三',))
        self.assertEqual(zhang['id_number'], '智能新身份证')
        self.assertEqual(zhang['bank_account'], '000456')
        self.assertEqual(len(db.query_all('SELECT * FROM employees')), 3)


if __name__ == '__main__':
    unittest.main()
