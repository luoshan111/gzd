"""XLS template fidelity and preview/confirm regression tests (no real employee DB)."""

import io
import os
import struct
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import xlwt
from flask import Flask
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import COLOR_INDEX

from core import config
from routes import import_export_routes
from services import import_export_service as exports, legacy_template, smart_fill_service as fill


def legacy_payroll():
    book = xlwt.Workbook(encoding='utf-8')
    sheet = book.add_sheet('工资表')
    style = xlwt.easyxf('font: name Arial, bold on; pattern: pattern solid, fore_colour yellow; '
                        'borders: left thin, right thin, top thin, bottom thin;', num_format_str='@')
    sheet.write_merge(0, 0, 0, 4, '项目工资表', style)
    for col, header in enumerate(['姓名', '身份证号', '银行卡号', '工资', '计算金额']):
        sheet.write(1, col, header, style)
    for col, value in enumerate(['旧人员', '旧证件', '000000', 300, xlwt.Formula('D3*2')]):
        sheet.write(2, col, value, style)
    sheet.write(3, 0, '合计')
    sheet.write(3, 4, xlwt.Formula('SUM(E3:E3)'))
    sheet.col(0).width = 20 * 256
    sheet.col(2).hidden = True
    sheet.row(2).height = 480
    sheet.row(2).height_mismatch = True
    sheet.panes_frozen = True
    sheet.horz_split_pos = 2
    sheet.portrait = False
    sheet.header_str = b'&CPayroll'
    sheet.footer_str = b'&CPage &P'
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def converted_payroll():
    book = Workbook()
    sheet = book.active
    sheet.title = '工资表'
    sheet.append(['项目工资表'])
    sheet.merge_cells('A1:E1')
    sheet.append(['姓名', '身份证号', '银行卡号', '工资', '计算金额'])
    sheet.append(['旧人员', '旧证件', '000000', 300, '=D3*2'])
    sheet.append(['合计', None, None, None, '=SUM(E3:E3)'])
    sheet['A3'].fill = PatternFill('solid', fgColor='FFFF00')
    sheet.column_dimensions['A'].width = 20
    sheet.row_dimensions[3].height = 24
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class LegacyTemplateTests(unittest.TestCase):
    def test_real_biff_passes_preflight_without_changing_bytes(self):
        source = legacy_payroll()
        original = bytes(source)
        legacy_template.validate_legacy_template(source)
        self.assertEqual(source, original)

    def test_invalid_and_oversized_files_are_rejected_before_excel(self):
        for data in (b'not-xls', legacy_template.OLE_MAGIC + b'0' * legacy_template.MAX_TEMPLATE_BYTES):
            with self.assertRaises(legacy_template.LegacyTemplateError):
                legacy_template.validate_legacy_template(data)

    def test_biff_macros_encryption_and_broken_records_are_rejected(self):
        for record, payload in ((0x002F, b'password'),
                                (0x0085, b'\x00' * 5 + b'\x01'),
                                (0x0085, b'\x00' * 5 + b'\x06'),
                                (0x0809, struct.pack('<HH', 0x0600, 0x0040))):
            with self.assertRaises(legacy_template.LegacyTemplateError):
                legacy_template._validate_biff_stream(struct.pack('<HH', record, len(payload)) + payload)
        with self.assertRaisesRegex(legacy_template.LegacyTemplateError, '损坏'):
            legacy_template._validate_biff_stream(struct.pack('<HH', 0x0085, 10) + b'0')

    def test_external_data_and_macro_project_markers_are_refused_before_open(self):
        for record in (0x00D3, 0x01AD, 0x0876, 0x00DC, 0x0803, 0x080A):
            with self.assertRaises(legacy_template.LegacyTemplateError):
                legacy_template._validate_biff_stream(struct.pack('<HH', record, 0))
        for external_kind in (0x3A01, 10):
            data = struct.pack('<HH', 1, external_kind)
            with self.assertRaisesRegex(legacy_template.LegacyTemplateError, '外部引用'):
                legacy_template._validate_biff_stream(struct.pack('<HH', 0x01AE, len(data)) + data)
        internal = struct.pack('<HH', 1, 0x0401)
        legacy_template._validate_biff_stream(struct.pack('<HH', 0x01AE, len(internal)) + internal)

    def test_unavailable_excel_has_actionable_error(self):
        with patch.object(legacy_template.subprocess, 'run',
                          return_value=subprocess.CompletedProcess([], 3)):
            with self.assertRaisesRegex(legacy_template.LegacyTemplateError, 'Microsoft Excel'):
                legacy_template._run_conversion(Path('source.xls'), Path('out.xlsx'), Path('owner.json'))

    def test_timeout_runs_only_owned_instance_cleanup(self):
        with patch.object(legacy_template.subprocess, 'run', side_effect=[
            subprocess.TimeoutExpired('powershell', 60), subprocess.CompletedProcess([], 0),
        ]) as run:
            with self.assertRaisesRegex(legacy_template.LegacyTemplateError, '超时'):
                legacy_template._run_conversion(Path('source.xls'), Path('out.xlsx'), Path('owner.json'))
        self.assertEqual(run.call_count, 2)
        self.assertEqual(run.call_args.args[0][-1], '-CleanupOnly')
        self.assertNotIn('/IM', run.call_args.args[0])

    @unittest.skipUnless(os.name == 'nt', 'Windows conversion worker')
    def test_temp_files_are_removed_after_conversion_failure(self):
        folders = []
        def fail(source, target, owner):
            folders.append(source.parent)
            self.assertEqual(source.read_bytes(), legacy_payroll())
            raise legacy_template.LegacyTemplateError('test failure')
        with patch.object(legacy_template, '_run_conversion', side_effect=fail):
            with self.assertRaises(legacy_template.LegacyTemplateError):
                legacy_template.convert_legacy_template(legacy_payroll())
        self.assertFalse(folders[0].exists())


class LegacyFillRouteTests(unittest.TestCase):
    def setUp(self):
        exports._smart_fill_preview_store.clear()
        self.addCleanup(exports._smart_fill_preview_store.clear)
        app = Flask(__name__)
        app.secret_key = 'xls-fill-test'
        import_export_routes.register_import_export_routes(app)
        self.client = app.test_client()
        with self.client.session_transaction() as session:
            session['user_id'] = 1
            session['username'] = 'test'
        key = patch.object(config, 'AGENT_API_KEY', '')
        key.start()
        self.addCleanup(key.stop)
        records = patch.object(fill, '_employee_records', return_value={
            '张三': {'real_name': '张三', 'id_number': '110101199001011234',
                     'bank_account': '000123456789', 'phone': '', 'bank_address': ''},
        })
        records.start()
        self.addCleanup(records.stop)

    def preview(self, source):
        return self.client.post('/api/smart-fill-preview', data={
            'file': (io.BytesIO(source), '客户端工资表.XLS'), 'names': '["张三"]', 'overwrite': '1',
        }, content_type='multipart/form-data')

    def confirm(self, token, source):
        return self.client.post('/api/smart-fill-confirm', data={
            'preview_token': token, 'file': (io.BytesIO(source), '客户端工资表.XLS'),
        }, content_type='multipart/form-data')

    def test_preview_converts_once_and_confirm_reuses_validated_workbook(self):
        source = legacy_payroll()
        with patch.object(fill, 'convert_legacy_template', return_value=converted_payroll()) as convert:
            preview = self.preview(source)
            self.assertEqual(preview.status_code, 200, preview.get_json())
            data = preview.get_json()
            self.assertTrue(data['preview']['converted'])
            self.assertEqual(data['preview']['source_format'], 'xls')
            self.assertEqual(data['preview']['output_format'], 'xlsx')
            token = data['preview_token']
            self.assertEqual(self.confirm(token, source + b'changed').status_code, 409)
            result = self.confirm(token, source)
            self.assertEqual(result.status_code, 200, result.get_json(silent=True))
            convert.assert_called_once_with(source)
        self.assertNotIn(token, exports._smart_fill_preview_store)
        self.assertIn('.xlsx', result.headers['Content-Disposition'])
        book = load_workbook(io.BytesIO(result.data))
        sheet = book['工资表']
        self.assertEqual(sheet['A3'].value, '张三')
        self.assertEqual(sheet['B3'].value, '110101199001011234')
        self.assertEqual(sheet['C3'].value, '000123456789')
        self.assertIsNone(sheet['D3'].value)
        self.assertEqual(sheet['E3'].value, '=D3*2')
        self.assertEqual(sheet['E4'].value, '=SUM(E3:E3)')
        self.assertEqual(str(sheet.merged_cells), 'A1:E1')
        self.assertEqual(sheet['A3'].fill.fgColor.rgb, '00FFFF00')
        self.assertEqual(sheet.row_dimensions[3].height, 24)
        self.assertEqual(sheet.column_dimensions['A'].width, 20)
        book.close()

    def test_conversion_error_does_not_issue_preview_or_modify_db(self):
        with patch.object(fill, 'convert_legacy_template',
                          side_effect=legacy_template.LegacyTemplateError('转换不可用')), \
             patch.object(fill, '_employee_records') as records:
            response = self.preview(legacy_payroll())
        self.assertEqual(response.status_code, 422)
        self.assertTrue(response.get_json()['fallback_available'])
        records.assert_not_called()
        self.assertFalse(exports._smart_fill_preview_store)

    def test_fill_error_is_422_not_file_conflict_409(self):
        with patch.object(fill, 'convert_legacy_template', return_value=converted_payroll()):
            preview = self.preview(legacy_payroll()).get_json()
        with patch.object(import_export_routes, 'fill_template', side_effect=fill.TemplateFillError('填写失败')):
            response = self.confirm(preview['preview_token'], legacy_payroll())
        self.assertEqual(response.status_code, 422)
        self.assertTrue(response.get_json()['fallback_available'])


@unittest.skipUnless(os.environ.get('GZD_TEST_EXCEL_CONVERSION') == '1',
                     'Opt-in: requires installed Microsoft Excel on Windows')
class NativeExcelConversionTests(unittest.TestCase):
    setUp = LegacyFillRouteTests.setUp
    preview = LegacyFillRouteTests.preview
    confirm = LegacyFillRouteTests.confirm

    def test_native_conversion_preserves_layout_and_formulas_and_fills(self):
        source = legacy_payroll()
        with tempfile.TemporaryDirectory(prefix='gzd-native-check-') as directory:
            original = Path(directory) / 'original.xls'
            original.write_bytes(source)
            converted = legacy_template.convert_legacy_template(original.read_bytes())
            self.assertEqual(original.read_bytes(), source)
        book = load_workbook(io.BytesIO(converted))
        sheet = book['工资表']
        self.assertEqual(sheet['E3'].value, '=D3*2')
        self.assertEqual(sheet['E4'].value, '=SUM(E3:E3)')
        self.assertEqual(str(sheet.merged_cells), 'A1:E1')
        color = sheet['A3'].fill.fgColor
        self.assertEqual((COLOR_INDEX[color.indexed] if color.type == 'indexed' else color.rgb)[-6:],
                         'FFFF00')
        self.assertEqual(sheet['A3'].border.left.style, 'thin')
        self.assertAlmostEqual(sheet.row_dimensions[3].height, 24)
        self.assertAlmostEqual(sheet.column_dimensions['A'].width, 20, delta=1)
        self.assertTrue(sheet.column_dimensions['C'].hidden)
        self.assertEqual(sheet.freeze_panes, 'A3')
        self.assertEqual(sheet.page_setup.orientation, 'landscape')
        self.assertEqual(sheet.oddHeader.center.text, 'Payroll')
        book.close()
        # Actual native conversion in the route too: no mocked converter or AI network.
        with patch.object(fill, 'convert_legacy_template', wraps=legacy_template.convert_legacy_template) as convert:
            preview = self.preview(source)
            self.assertEqual(preview.status_code, 200, preview.get_json())
            confirmed = self.confirm(preview.get_json()['preview_token'], source)
            self.assertEqual(confirmed.status_code, 200, confirmed.get_json(silent=True))
            convert.assert_called_once()
        book = load_workbook(io.BytesIO(confirmed.data))
        self.assertEqual(book['工资表']['A3'].value, '张三')
        self.assertEqual(book['工资表']['C3'].value, '000123456789')
        self.assertEqual(book['工资表']['E3'].value, '=D3*2')
        book.close()


if __name__ == '__main__':
    unittest.main()
